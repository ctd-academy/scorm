#!/usr/bin/env python3
"""Đóng gói 1 khóa học thành SCORM 1.2 từ file Excel nhập liệu + thư mục video.

Cách dùng:
  python3 tools/build_scorm.py <file_excel> <thư_mục_khóa_chứa_video> <thư_mục_xuất>

Đây là "động cơ" tham chiếu của Builder. Builder web (builder/) phải cho ra gói
tương đương: cùng cấu trúc zip, cùng course.js, cùng imsmanifest.xml.
"""
import json, os, re, sys, zipfile
from xml.sax.saxutils import escape
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
PLAYER = os.path.join(HERE, "..", "player")
LETTERS = "ABCD"


def cell(v):
    return "" if v is None else str(v).strip()


def parse_time(s):
    # Nếu Excel tự đổi "01:30" thành giờ:phút, hiểu lại là phút:giây như người nhập muốn
    if hasattr(s, "hour") and hasattr(s, "minute"):
        return s.hour * 60 + s.minute
    s = cell(s)
    m = re.fullmatch(r"(?:(\d+):)?(\d{1,3}):(\d{1,2})", s)
    if not m:
        return None
    h, mm, ss = int(m.group(1) or 0), int(m.group(2)), int(m.group(3))
    return h * 3600 + mm * 60 + ss


def video_duration(course_dir, fname):
    """Đọc thời lượng video (giây) để hiện trong mục lục và tính thời lượng khóa."""
    import subprocess
    for p in (os.path.join(course_dir, fname), os.path.join(course_dir, "videos", fname)):
        if os.path.isfile(p):
            try:
                out = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration",
                                      "-of", "default=nw=1:nk=1", p], capture_output=True, text=True, timeout=30)
                return round(float(out.stdout.strip()))
            except Exception:
                return 0
    return 0


def read_excel(path, course_dir):
    wb = load_workbook(path, data_only=True)
    errors = []

    info = {}
    for row in wb["Thông tin khóa"].iter_rows(min_row=2, max_row=11, values_only=True):
        if row[0]:
            info[cell(row[0])] = row[1]

    def num(key, default):
        try:
            return int(float(info.get(key)))
        except (TypeError, ValueError):
            return default

    course_id = re.sub(r"[^A-Za-z0-9_-]", "", cell(info.get("Mã khóa"))) or "KHOA"
    settings = {
        "passScore": num("Điểm đạt (%)", 80),
        "maxAttempts": num("Số lượt làm bài", 3),
        "questionsPerAttempt": num("Số câu mỗi lần thi", 0),
        "timeLimitMinutes": num("Thời gian làm bài (phút)", 0),
        "showAnswers": cell(info.get("Hiện đáp án sau khi nộp")).lower() in ("có", "co", "yes", "x"),
        "maxSpeed": 2,
        "contact": cell(info.get("Nội dung liên hệ khi hết lượt")) or "Bạn đã dùng hết lượt làm bài. Vui lòng liên hệ L&OD: Mr. Minh (minhvt02@coteccons.vn) hoặc cta@coteccons.vn để được hỗ trợ đăng ký lại.",
    }
    logo = cell(info.get("File logo (tùy chọn)"))

    # Video: dòng 2 là ví dụ, bỏ qua nếu cột STT ghi "(Ví dụ)"
    videos = []
    for r, row in enumerate(wb["Video"].iter_rows(min_row=2, values_only=True), 2):
        if not row or not cell(row[2]) or cell(row[0]).startswith("(Ví dụ)"):
            continue
        fname = cell(row[2])
        if not os.path.isfile(os.path.join(course_dir, fname)) and not os.path.isfile(os.path.join(course_dir, "videos", fname)):
            errors.append(f"Video dòng {r}: không tìm thấy file '{fname}'")
        videos.append({"title": cell(row[1]) or f"Bài {len(videos)+1}", "file": fname})

    popups, quiz = [], []
    for r, row in enumerate(wb["Câu hỏi"].iter_rows(min_row=2, values_only=True), 2):
        if not row or not cell(row[3]) or cell(row[3]).startswith("(Ví dụ)"):
            continue
        kind = cell(row[0]).upper()
        opts = [cell(x) for x in row[4:8]]
        while opts and not opts[-1]:
            opts.pop()
        if len(opts) < 2 or "" in opts:
            errors.append(f"Câu hỏi dòng {r}: cần ít nhất phương án A, B và không bỏ trống ở giữa")
            continue
        ans = cell(row[8]).upper()
        if ans not in LETTERS[:len(opts)]:
            errors.append(f"Câu hỏi dòng {r}: đáp án đúng '{ans}' không hợp lệ")
            continue
        q = {"q": cell(row[3]), "options": opts, "answer": LETTERS.index(ans), "explain": cell(row[9])}
        if kind == "POPUP":
            try:
                vno = int(float(row[1]))
            except (TypeError, ValueError):
                vno = 0
            t = parse_time(row[2])
            if not (1 <= vno <= len(videos)):
                errors.append(f"Câu hỏi dòng {r}: 'Video số' phải từ 1 đến {len(videos)}")
                continue
            if t is None:
                errors.append(f"Câu hỏi dòng {r}: mốc thời gian '{cell(row[2])}' phải có dạng mm:ss")
                continue
            q.update({"video": vno - 1, "time": t})
            popups.append(q)
        elif kind in ("CUOI KHOA", "CUỐI KHÓA"):
            quiz.append(q)
        else:
            errors.append(f"Câu hỏi dòng {r}: loại '{kind}' phải là POPUP hoặc CUOI KHOA")

    if not videos:
        errors.append("Chưa có video nào trong sheet 'Video'")
    if quiz and settings["questionsPerAttempt"] > len(quiz):
        errors.append(f"Số câu mỗi lần thi ({settings['questionsPerAttempt']}) nhiều hơn ngân hàng ({len(quiz)} câu)")

    course = {
        "id": course_id,
        "title": cell(info.get("Tên khóa")) or course_id,
        "description": cell(info.get("Mô tả ngắn")),
        "logo": ("assets/" + logo) if logo else None,
        "settings": settings,
        "videos": [{"title": v["title"], "src": "videos/" + v["file"],
                    "duration": video_duration(course_dir, v["file"])} for v in videos],
        "popups": popups,
        "quiz": quiz,
    }
    return course, logo, errors


def manifest(course, files):
    file_tags = "\n".join(f'      <file href="{escape(f)}"/>' for f in files)
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<manifest identifier="CTD_{escape(course['id'])}" version="1.0"
  xmlns="http://www.imsproject.org/xsd/imscp_rootv1p1p2"
  xmlns:adlcp="http://www.adlnet.org/xsd/adlcp_rootv1p2"
  xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance"
  xsi:schemaLocation="http://www.imsproject.org/xsd/imscp_rootv1p1p2 imscp_rootv1p1p2.xsd
                      http://www.imsglobal.org/xsd/imsmd_rootv1p2p1 imsmd_rootv1p2p1.xsd
                      http://www.adlnet.org/xsd/adlcp_rootv1p2 adlcp_rootv1p2.xsd">
  <metadata><schema>ADL SCORM</schema><schemaversion>1.2</schemaversion></metadata>
  <organizations default="ORG1">
    <organization identifier="ORG1">
      <title>{escape(course['title'])}</title>
      <item identifier="ITEM1" identifierref="RES1">
        <title>{escape(course['title'])}</title>
      </item>
    </organization>
  </organizations>
  <resources>
    <resource identifier="RES1" type="webcontent" adlcp:scormtype="sco" href="index.html">
{file_tags}
    </resource>
  </resources>
</manifest>
"""


def build(xlsx, course_dir, out_dir):
    course, logo, errors = read_excel(xlsx, course_dir)
    if errors:
        print("KHÔNG XUẤT ĐƯỢC. Hãy sửa các lỗi sau trong file Excel:")
        for e in errors:
            print("  -", e)
        return None

    entries = []  # (đường dẫn trong zip, nguồn hoặc nội dung bytes)
    for root, _, names in os.walk(PLAYER):
        for n in sorted(names):
            if n == "course.js":
                continue
            full = os.path.join(root, n)
            entries.append((os.path.relpath(full, PLAYER).replace(os.sep, "/"), full))
    js = "window.COURSE = " + json.dumps(course, ensure_ascii=False, indent=1) + ";\n"
    entries.append(("course.js", js.encode("utf-8")))
    for v in course["videos"]:
        fname = v["src"][len("videos/"):]
        src = os.path.join(course_dir, fname)
        if not os.path.isfile(src):
            src = os.path.join(course_dir, "videos", fname)
        entries.append((v["src"], src))
    if logo:
        entries.append(("assets/" + logo, os.path.join(course_dir, logo)))
    files = [e[0] for e in entries]
    entries.append(("imsmanifest.xml", manifest(course, files).encode("utf-8")))

    os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, f"SCORM_{course['id']}.zip")
    with zipfile.ZipFile(out, "w") as z:
        for name, src in entries:
            # video đã nén sẵn -> lưu nguyên (STORED) cho nhanh; còn lại nén DEFLATED
            ctype = zipfile.ZIP_STORED if name.endswith((".mp4", ".woff2", ".png", ".jpg")) else zipfile.ZIP_DEFLATED
            if isinstance(src, bytes):
                z.writestr(zipfile.ZipInfo(name, (2026, 1, 1, 0, 0, 0)), src, compress_type=ctype)
            else:
                z.write(src, name, compress_type=ctype)
    size = os.path.getsize(out) / 1024 / 1024
    print(f"Đã xuất: {out} ({size:.1f} MB)")
    total = sum(v["duration"] for v in course["videos"])
    print(f"  {len(course['videos'])} video · {len(course['popups'])} pop-up · {len(course['quiz'])} câu ngân hàng đề")
    print(f"  Tổng thời lượng video: {total // 60} phút {total % 60} giây (dùng cho trường Total Course Duration trên Workday)")
    return out, course


if __name__ == "__main__":
    if len(sys.argv) != 4:
        print(__doc__)
        sys.exit(1)
    sys.exit(0 if build(*sys.argv[1:]) else 2)
