#!/usr/bin/env python3
"""Tạo file Excel nhập liệu mẫu (kèm dữ liệu khóa demo ATLD-DEMO).

Cách dùng:  python3 tools/make_template.py [đường_dẫn_xuất]
Mặc định:   templates/Mau_nhap_lieu_khoa_hoc.xlsx

LƯU Ý: Alex dùng Excel bản cũ, công thức trong file mở bằng tay có thể cần dấu ';'.
openpyxl luôn ghi công thức chuẩn dấu ',' và Excel tự hiển thị theo locale máy.
"""
import os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "..", "templates", "Mau_nhap_lieu_khoa_hoc.xlsx")

NAVY="16315E"; TEAL="5FD1C1"; FILL="FFF7D6"
F="Arial"
hdr_font=Font(name=F,bold=True,color="FFFFFF",size=11)
hdr_fill=PatternFill("solid",fgColor=NAVY)
inp=PatternFill("solid",fgColor=FILL)
thin=Side(style="thin",color="C4CAD6"); box=Border(left=thin,right=thin,top=thin,bottom=thin)
body=Font(name=F,size=10); bold=Font(name=F,size=10,bold=True)
wrap=Alignment(wrap_text=True,vertical="top")

def header(ws,row,cols):
    for i,c in enumerate(cols,1):
        cell=ws.cell(row=row,column=i,value=c); cell.font=hdr_font; cell.fill=hdr_fill
        cell.alignment=Alignment(wrap_text=True,vertical="center"); cell.border=box

wb=Workbook()
# ---------------- Hướng dẫn
g=wb.active; g.title="Hướng dẫn"
g["A1"]="FILE NHẬP LIỆU KHÓA HỌC E-LEARNING (SCORM) – COTECCONS L&OD"; g["A1"].font=Font(name=F,size=14,bold=True,color=NAVY)
lines=[
("Cách dùng",""),
("1","Mỗi khóa học = 1 file Excel này. Đặt tên file theo mã khóa, ví dụ: ATLD-001.xlsx"),
("2","Ô nền vàng nhạt là ô cần điền. Không đổi tên sheet, không đổi tiêu đề cột."),
("3","Sheet 'Thông tin khóa': thông tin chung và luật bài kiểm tra của khóa."),
("4","Sheet 'Video': mỗi dòng là 1 bài (1 video), theo đúng thứ tự học. Tên file phải trùng tên file video trong thư mục khóa."),
("5","Sheet 'Câu hỏi': gồm cả câu hỏi pop-up (giữa video) và câu hỏi bài kiểm tra cuối khóa."),
("6","Dòng 2 ở mỗi sheet (chữ nghiêng) là VÍ DỤ. Xóa hoặc ghi đè khi nhập khóa thật."),
("",""),
("Luật chơi cố định",""),
("•","Học tuần tự: bài sau chỉ mở khi xem xong bài trước."),
("•","Video chỉ tua lại phần đã xem, không tua tới. Tốc độ tối đa 2x. Rời trang/chuyển tab thì video tự dừng."),
("•","Pop-up không tính điểm. Trả lời sai thì phải xem lại đoạn video trước câu hỏi (từ pop-up trước đó, hoặc từ đầu video)."),
("•","Bài kiểm tra chỉ mở khi đã xem hết video. Khóa không có câu hỏi CUỐI KHÓA sẽ hoàn thành khi xem hết video."),
("•","Hết lượt làm bài: bài kiểm tra bị khóa, học viên thấy nội dung liên hệ L&OD. L&OD đăng ký lại cho học viên trên Workday."),
("",""),
("Quy ước nhập câu hỏi",""),
("Loại","POPUP = câu hỏi giữa video · CUOI KHOA = câu hỏi bài kiểm tra cuối khóa"),
("Video số","Chỉ điền cho POPUP: số thứ tự bài (cột STT ở sheet Video)"),
("Mốc thời gian","Chỉ điền cho POPUP: phút:giây trong video, ví dụ 01:30. Nên đặt cách nhau ít nhất 1 phút"),
("Phương án","Tối thiểu A, B. C và D có thể bỏ trống"),
("Đáp án đúng","Một chữ cái: A, B, C hoặc D"),
("Giải thích","Không bắt buộc. Hiện sau khi trả lời pop-up đúng, và khi xem lại bài kiểm tra (nếu bật hiện đáp án)"),
]
for i,(a,b) in enumerate(lines,3):
    g.cell(row=i,column=1,value=a).font=bold if b=="" or a in("Loại","Video số","Mốc thời gian","Phương án","Đáp án đúng","Giải thích") else body
    if b=="" and a: g.cell(row=i,column=1).font=Font(name=F,size=11,bold=True,color=NAVY)
    c=g.cell(row=i,column=2,value=b); c.font=body; c.alignment=wrap
g.column_dimensions["A"].width=16; g.column_dimensions["B"].width=110
g.cell(row=3+len(lines)+1,column=1,value="Ô cần điền").font=bold
g.cell(row=3+len(lines)+1,column=2).fill=inp; g.cell(row=3+len(lines)+1,column=2).border=box

# ---------------- Thông tin khóa
t=wb.create_sheet("Thông tin khóa")
header(t,1,["Mục","Giá trị","Ghi chú"])
rows=[
("Mã khóa","ATLD-DEMO","Không dấu, không khoảng trắng. Dùng làm tên file SCORM"),
("Tên khóa","An toàn lao động cơ bản (Demo)","Hiển thị trên đầu trang khóa học"),
("Mô tả ngắn","Khóa demo minh họa các tính năng của Coteccons SCORM Player. Nội dung chỉ mang tính minh họa.","Dùng khi tạo khóa trên Workday"),
("Điểm đạt (%)",80,"Từ 0 đến 100"),
("Số lượt làm bài",3,"Số nguyên ≥ 1"),
("Số câu mỗi lần thi",5,"Rút ngẫu nhiên từ ngân hàng câu hỏi. 0 = dùng tất cả"),
("Thời gian làm bài (phút)",10,"0 = không giới hạn"),
("Hiện đáp án sau khi nộp","Có","Có / Không"),
("Nội dung liên hệ khi hết lượt","Bạn đã dùng hết lượt làm bài. Vui lòng liên hệ L&OD Team để được hỗ trợ đăng ký lại. (Thay bằng tên, email, số điện thoại đầu mối L&OD)","Học viên thấy nội dung này khi hết lượt"),
("File logo (tùy chọn)","","Tên file logo đặt trong thư mục khóa. Để trống = dùng chữ COTECCONS"),
]
for i,(a,b,c) in enumerate(rows,2):
    t.cell(row=i,column=1,value=a).font=bold
    v=t.cell(row=i,column=2,value=b); v.font=body; v.fill=inp; v.border=box; v.alignment=wrap
    n=t.cell(row=i,column=3,value=c); n.font=Font(name=F,size=9,italic=True,color="636468"); n.alignment=wrap
r=len(rows)+3
t.cell(row=r,column=1,value="Tự động tính (không sửa)").font=Font(name=F,size=10,bold=True,color=NAVY)
auto=[("Số video","=COUNTA(Video!C3:C500)"),
      ("Số câu pop-up","=COUNTIF('Câu hỏi'!A3:A2000,\"POPUP\")"),
      ("Số câu trong ngân hàng đề","=COUNTIF('Câu hỏi'!A3:A2000,\"CUOI KHOA\")"),
      ("Kiểm tra số câu mỗi lần thi","=IF(B7>B{bank},\"Lỗi: nhiều hơn số câu trong ngân hàng\",\"OK\")")]
bankrow=r+3
for k,(a,f) in enumerate(auto,1):
    t.cell(row=r+k,column=1,value=a).font=body
    t.cell(row=r+k,column=2,value=f.replace("{bank}",str(bankrow))).font=body
t.cell(row=r+5,column=1,value="Lưu ý: số liệu tự động chỉ đếm từ dòng 3 trở đi (dòng 2 là ví dụ ở sheet Video và Câu hỏi).").font=Font(name=F,size=9,italic=True,color="636468")
t.column_dimensions["A"].width=30; t.column_dimensions["B"].width=60; t.column_dimensions["C"].width=50
dv=DataValidation(type="list",formula1='"Có,Không"',allow_blank=False); t.add_data_validation(dv); dv.add("B9")
for addr,lo,hi in (("B5",0,100),("B6",1,20),("B7",0,500),("B8",0,600)):
    d=DataValidation(type="whole",operator="between",formula1=str(lo),formula2=str(hi),showErrorMessage=True,errorTitle="Giá trị không hợp lệ",error=f"Nhập số nguyên từ {lo} đến {hi}")
    t.add_data_validation(d); d.add(addr)

# ---------------- Video
v=wb.create_sheet("Video")
header(v,1,["STT","Tên bài","Tên file video","Ghi chú"])
v["A2"]="(Ví dụ)"; v["B2"]="Giới thiệu khóa học"; v["C2"]="01_gioi-thieu.mp4"; v["D2"]="Dòng ví dụ – nhập dữ liệu thật từ dòng 3"
for c in "ABCD": v[f"{c}2"].font=Font(name=F,size=10,italic=True,color="636468")
vids=[(1,"Vì sao an toàn là trên hết","01_vi-sao-an-toan.mp4",""),
      (2,"Trang bị bảo hộ cá nhân (PPE)","02_bao-ho-ca-nhan.mp4",""),
      (3,"Nhận diện mối nguy tại công trường","03_nhan-dien-moi-nguy.mp4","")]
for i,row in enumerate(vids,3):
    for j,val in enumerate(row,1):
        c=v.cell(row=i,column=j,value=val); c.font=body; c.fill=inp; c.border=box
for i in range(len(vids)+3,60):
    for j in range(1,5): v.cell(row=i,column=j).fill=inp; v.cell(row=i,column=j).border=box
v.column_dimensions["A"].width=7; v.column_dimensions["B"].width=45; v.column_dimensions["C"].width=34; v.column_dimensions["D"].width=40
v.freeze_panes="A3"

# ---------------- Câu hỏi
q=wb.create_sheet("Câu hỏi")
cols=["Loại","Video số","Mốc thời gian (mm:ss)","Câu hỏi","A","B","C","D","Đáp án đúng","Giải thích"]
header(q,1,cols)
ex=["POPUP",1,"00:20","(Ví dụ) Câu hỏi pop-up ở giây 20 của video 1?","Phương án A","Phương án B","","","A","Dòng ví dụ – nhập dữ liệu thật từ dòng 3"]
for j,val in enumerate(ex,1):
    c=q.cell(row=2,column=j,value=val); c.font=Font(name=F,size=10,italic=True,color="636468")
data=[
["POPUP",1,"00:15","Theo video, điều gì được đặt lên hàng đầu tại công trường?","Tiến độ","An toàn","Chi phí","","B","An toàn luôn được ưu tiên trước tiến độ và chi phí."],
["POPUP",1,"00:40","Ai có trách nhiệm về an toàn tại công trường?","Chỉ cán bộ an toàn","Chỉ chỉ huy trưởng","Tất cả mọi người","","C","An toàn là trách nhiệm của mỗi người."],
["POPUP",2,"00:25","Mũ bảo hộ cần được đội như thế nào?","Đội lỏng cho thoải mái","Cài quai cằm chắc chắn","Chỉ đội khi có người kiểm tra","","B",""],
["POPUP",3,"00:30","Khi phát hiện mối nguy, việc đầu tiên nên làm là gì?","Tiếp tục làm việc","Báo ngay cho người phụ trách","Tự xử lý một mình","","B",""],
["CUOI KHOA","","","Mục đích chính của trang bị bảo hộ cá nhân là gì?","Làm đẹp đồng phục","Giảm nguy cơ chấn thương","Thay thế biện pháp quản lý","","B","PPE là lớp bảo vệ cuối cùng, giúp giảm nguy cơ chấn thương."],
["CUOI KHOA","","","Khi thấy đồng nghiệp không đội mũ bảo hộ, bạn nên làm gì?","Im lặng","Nhắc nhở ngay","Chụp ảnh đăng mạng xã hội","Báo cáo sau một tuần","B",""],
["CUOI KHOA","","","Phát hiện dây điện bị hở trên lối đi, bạn nên làm gì trước tiên?","Bước qua và đi tiếp","Rào chắn khu vực và báo người phụ trách","Tự nối lại dây điện","","B",""],
["CUOI KHOA","","","An toàn lao động là trách nhiệm của ai?","Riêng bộ phận an toàn","Riêng ban chỉ huy","Mọi người trên công trường","","C",""],
["CUOI KHOA","","","Trước khi bắt đầu ca làm việc, bạn nên làm gì?","Kiểm tra trang bị bảo hộ và khu vực làm việc","Bắt đầu ngay để kịp tiến độ","Chờ có người nhắc","","A",""],
["CUOI KHOA","","","Khi không chắc chắn về quy trình an toàn của một công việc, bạn nên?","Làm theo cách mình nghĩ","Hỏi người phụ trách trước khi làm","Bỏ qua bước không rõ","","B",""],
]
for i,row in enumerate(data,3):
    for j,val in enumerate(row,1):
        c=q.cell(row=i,column=j,value=val); c.font=body; c.fill=inp; c.border=box; c.alignment=wrap
for i in range(len(data)+3,300):
    for j in range(1,11): q.cell(row=i,column=j).fill=inp; q.cell(row=i,column=j).border=box
for i in range(2,300): q.cell(row=i,column=3).number_format="@"
widths=[12,9,13,50,26,26,26,26,10,40]
for j,w in enumerate(widths,1): q.column_dimensions[chr(64+j)].width=w
q.freeze_panes="A3"
d1=DataValidation(type="list",formula1='"POPUP,CUOI KHOA"',allow_blank=True); q.add_data_validation(d1); d1.add("A3:A2000")
d2=DataValidation(type="list",formula1='"A,B,C,D"',allow_blank=True); q.add_data_validation(d2); d2.add("I3:I2000")
q["C1"].comment=Comment("Định dạng phút:giây, ví dụ 01:30. Chỉ điền cho POPUP.","L&OD")

os.makedirs(os.path.dirname(os.path.abspath(OUT)), exist_ok=True)
wb.save(OUT)
print("Đã tạo", OUT)
